import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
import csv



def convert_excel():
    file = filedialog.askopenfilename(
        title="Select Excel File",
        filetypes=[("Excel Files", "*.xlsx")]
    )

    if not file:
        return

    try:
        workbook = load_workbook(file, data_only=True)
        sheet = workbook.active

        save = filedialog.asksaveasfilename(
            title="Save CSV File",
            defaultextension=".csv",
            filetypes=[("CSV Files", "*.csv")]
        )

        if not save:
            return

        with open(save, "w", newline="", encoding="utf-8-sig") as csv_file:
            writer = csv.writer(csv_file)

            for row in sheet.iter_rows(values_only=True):
                writer.writerow(row)

        messagebox.showinfo(
            "Success ✨",
            "Excel file converted to CSV successfully!"
        )

        status.config(text="✓ Conversion completed!")

    except Exception as e:
        messagebox.showerror("Error", str(e))


# ---------- MAIN WINDOW ----------

root = tk.Tk()
root.title("Excel to CSV Converter ♡")
root.geometry("550x430")
root.resizable(False, False)
root.configure(bg="#FFF7F0")


# ---------- HEADER ----------

header = tk.Frame(root, bg="#5B4B8A", height=120)
header.pack(fill="x")

tk.Label(
    header,
    text="📊 Excel → CSV",
    font=("Georgia", 26, "bold"),
    bg="#5B4B8A",
    fg="white"
).pack(pady=(25, 5))

tk.Label(
    header,
    text="Simple • Cute • Easy Converter",
    font=("Arial", 11),
    bg="#5B4B8A",
    fg="#EDE7F6"
).pack()


# ---------- MAIN ----------

tk.Label(
    root,
    text="Convert your Excel file into CSV ✨",
    font=("Georgia", 17, "bold"),
    bg="#FFF7F0",
    fg="#3F365C"
).pack(pady=(35, 20))


tk.Button(
    root,
    text="📂  Select Excel File",
    command=convert_excel,
    font=("Arial", 14, "bold"),
    bg="#8E7CC3",
    fg="white",
    activebackground="#7565A5",
    relief="flat",
    width=25,
    height=2,
    cursor="hand2"
).pack(pady=15)


status = tk.Label(
    root,
    text="Ready to convert ✨",
    font=("Arial", 10),
    bg="#FFF7F0",
    fg="#8A7F8D"
)
status.pack(pady=20)


# ---------- FOOTER ----------

tk.Label(
    root,
    font=("Arial", 10),
    bg="#FFF7F0",
    fg="#8A7F8D"
).pack(pady=(30, 3))

tk.Label(
    root,
    font=("Georgia", 11, "bold"),
    bg="#FFF7F0",
    fg="#5B4B8A"
).pack()


root.mainloop()